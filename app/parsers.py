"""把上传文件转成纯文本。支持 txt/md/csv/docx/pdf/pptx/xlsx。"""
from __future__ import annotations

import io
from pathlib import Path

SUPPORTED = {".txt", ".md", ".csv", ".docx", ".pdf", ".pptx", ".xlsx"}


class ParseError(Exception):
    pass


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _docx(data: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    parts: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") or style.startswith("标题"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            parts.append("#" * min(int(level), 4) + " " + text)
        else:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                parts.append("|".join(cells))
    return "\n\n".join(parts)


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[第{i}页]\n{text}")
    return "\n\n".join(pages)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    texts.append("|".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append("备注:" + slide.notes_slide.notes_text_frame.text.strip())
        if texts:
            slides.append(f"[第{i}张幻灯片]\n" + "\n".join(texts))
    return "\n\n".join(slides)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                rows.append("|".join(cells))
        if rows:
            parts.append(f"## 工作表:{ws.title}\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        hint = "(老版 .doc/.ppt/.xls 请先另存为 .docx/.pptx/.xlsx)" if ext in {".doc", ".ppt", ".xls"} else ""
        raise ParseError(f"暂不支持 {ext or '无扩展名'} 文件{hint}")
    try:
        if ext in {".txt", ".md", ".csv"}:
            text = _decode(data)
        elif ext == ".docx":
            text = _docx(data)
        elif ext == ".pdf":
            text = _pdf(data)
        elif ext == ".pptx":
            text = _pptx(data)
        else:
            text = _xlsx(data)
    except ParseError:
        raise
    except Exception as exc:  # 文件损坏、加密等
        raise ParseError(f"文件解析失败:{exc}") from exc
    text = text.replace("\r\n", "\n").replace("\x00", "").strip()
    if not text:
        if ext == ".pdf":
            raise ParseError("PDF 中没有可提取的文字,可能是扫描版。请先用 OCR 工具转成文字版再上传。")
        raise ParseError("文件中没有可提取的文字")
    return text
