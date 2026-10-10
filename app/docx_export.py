"""把助手的回答(教案、培训方案、话术等)转成 Word 文档,老师下载后可直接打印或修改。"""
from __future__ import annotations

import io
import re
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "微软雅黑"
_BOLD = re.compile(r"\*\*([^*]+)\*\*")


def _set_font(run, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _add_inline(paragraph, text: str, size: float = 11) -> None:
    """支持 **加粗**,去掉行内代码的反引号。"""
    text = text.replace("`", "")
    pos = 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            _set_font(paragraph.add_run(text[pos:m.start()]), size)
        _set_font(paragraph.add_run(m.group(1)), size, bold=True)
        pos = m.end()
    if pos < len(text):
        _set_font(paragraph.add_run(text[pos:]), size)


def safe_filename(title: str) -> str:
    name = re.sub(r"[\\/:*?\"<>|\r\n]+", "_", title).strip(" ._")[:40]
    return name or f"教研助手_{date.today():%Y%m%d}"


def build_docx(content: str, title: str = "") -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)

    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(p.add_run(title), 16, bold=True)

    in_code = False
    for raw in content.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph()
            _set_font(p.add_run(line), 10, color=RGBColor(0x44, 0x44, 0x44))
            continue
        if not line.strip():
            continue
        if re.match(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$", line):
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = min(len(m.group(1)), 3)
            p = doc.add_paragraph()
            _set_font(p.add_run(m.group(2).replace("**", "")), {1: 15, 2: 13.5, 3: 12}[level], bold=True)
            continue
        m = re.match(r"^\s*[-*•]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, m.group(1))
            continue
        m = re.match(r"^>\s?(.*)$", line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _set_font(p.add_run(m.group(1).replace("**", "")), 10.5, color=RGBColor(0x66, 0x66, 0x66))
            continue
        # 中文编号("一、""1.""(1)")保留原样,便于老师按格式修改
        p = doc.add_paragraph()
        if re.match(r"^[一二三四五六七八九十]+、", line.strip()):
            _set_font(p.add_run(line.strip().replace("**", "")), 12, bold=True)
        else:
            _add_inline(p, line.strip())

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(footer.add_run(f"优学东方内部教研资料 · 仅供内部教学使用 · {date.today():%Y-%m-%d}"), 9,
              color=RGBColor(0x88, 0x88, 0x88))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
